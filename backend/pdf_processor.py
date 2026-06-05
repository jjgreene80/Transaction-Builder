"""
PDF and image processing: text extraction, OCR detection, packet splitting,
image-to-PDF conversion.
"""

import io
import os
import re
from pathlib import Path

import fitz  # PyMuPDF
import pdfplumber
from PIL import Image

# Optional OCR — gracefully disabled if Tesseract is not installed
try:
    import pytesseract
    _OCR_AVAILABLE = True
except ImportError:
    _OCR_AVAILABLE = False

# Text volume threshold below which a page is considered "scanned / image-only"
_TEXT_THRESHOLD = 50


# ---------------------------------------------------------------------------
# Public helpers
# ---------------------------------------------------------------------------

def get_page_count(pdf_path: str) -> int:
    try:
        doc = fitz.open(pdf_path)
        n = len(doc)
        doc.close()
        return n
    except Exception:
        return 1


def needs_ocr(pdf_path: str) -> bool:
    """Return True if the majority of pages lack extractable text."""
    if not _OCR_AVAILABLE:
        return False
    try:
        doc = fitz.open(pdf_path)
        if len(doc) == 0:
            doc.close()
            return False
        low_text_pages = sum(
            1 for page in doc if len(page.get_text().strip()) < _TEXT_THRESHOLD
        )
        doc.close()
        return low_text_pages > len(fitz.open(pdf_path)) / 2
    except Exception:
        return False


def extract_text(file_path: str) -> str:
    """Extract plain text from PDF, DOCX, XLSX, or image file."""
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        return _extract_pdf_text(file_path)
    elif ext == ".docx":
        return _extract_docx_text(file_path)
    elif ext == ".xlsx":
        return _extract_xlsx_text(file_path)
    elif ext in {".jpg", ".jpeg", ".png", ".tiff", ".tif"}:
        return _ocr_image_file(file_path)
    return ""


def ocr_pdf(input_path: str, output_path: str) -> None:
    """
    Create a copy of a scanned PDF with OCR text extracted and written
    as invisible text on each page so the file becomes searchable.
    Falls back to copying the original if OCR is unavailable.
    """
    if not _OCR_AVAILABLE:
        import shutil
        shutil.copy2(input_path, output_path)
        return

    src = fitz.open(input_path)
    out = fitz.open()

    for page_num in range(len(src)):
        src_page = src[page_num]
        has_text = len(src_page.get_text().strip()) >= _TEXT_THRESHOLD

        if has_text:
            out.insert_pdf(src, from_page=page_num, to_page=page_num)
            continue

        # Render page to image for OCR
        mat = fitz.Matrix(2.0, 2.0)  # 144 DPI
        pix = src_page.get_pixmap(matrix=mat)
        img_bytes = pix.tobytes("png")

        ocr_text = _ocr_image_bytes(img_bytes)

        # Insert original page then overlay invisible text
        out.insert_pdf(src, from_page=page_num, to_page=page_num)
        new_page = out[-1]
        if ocr_text.strip():
            new_page.insert_text(
                (10, 10),
                ocr_text,
                fontsize=1,
                color=(1, 1, 1),  # white = invisible
                render_mode=3,
            )

    out.save(output_path)
    out.close()
    src.close()


def split_pdf(pdf_path: str, output_dir: str) -> list[str]:
    """
    Attempt to split a multi-document packet PDF at document boundaries.
    Returns a list of output file paths (may be just [pdf_path] if no splits found).
    """
    doc = fitz.open(pdf_path)
    total_pages = len(doc)

    if total_pages <= 2:
        doc.close()
        return [pdf_path]

    # Extract first 400 chars of text per page
    page_texts = [doc[i].get_text()[:400] for i in range(total_pages)]

    # Find split points
    split_points = [0]
    for i in range(1, total_pages):
        if _is_new_doc_page(page_texts[i], page_texts[i - 1]):
            split_points.append(i)

    if len(split_points) == 1:
        doc.close()
        return [pdf_path]

    split_points.append(total_pages)
    output_paths = []
    stem = Path(pdf_path).stem

    for idx in range(len(split_points) - 1):
        start = split_points[idx]
        end = split_points[idx + 1] - 1
        new_doc = fitz.open()
        new_doc.insert_pdf(doc, from_page=start, to_page=end)
        out_path = os.path.join(output_dir, f"{stem}_part{idx + 1:02d}.pdf")
        new_doc.save(out_path)
        new_doc.close()
        output_paths.append(out_path)

    doc.close()
    return output_paths


def convert_image_to_pdf(image_path: str, output_pdf_path: str) -> None:
    """Convert a JPG/PNG/TIFF image to a single-page PDF."""
    img = Image.open(image_path).convert("RGB")
    img.save(output_pdf_path, format="PDF", resolution=150)


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

_DOC_STARTERS = [
    "california residential purchase",
    "purchase agreement",
    "transfer disclosure",
    "agency disclosure",
    "natural hazard",
    "inspection report",
    "appraisal report",
    "preliminary report",
    "escrow instructions",
    "addendum",
    "counter offer",
    "closing disclosure",
    "loan estimate",
    "seller property questionnaire",
    "statewide buyer",
    "lead-based paint",
    "pest control",
    "termite inspection",
]


def _is_new_doc_page(text: str, prev_text: str) -> bool:
    """Heuristic: is this page likely the first page of a new document?"""
    stripped = text.strip()
    prev_stripped = prev_text.strip()

    # Blank separator page followed by content
    if len(prev_stripped) < 80 and len(stripped) > 200:
        return True

    lower = stripped.lower()
    first_200 = lower[:200]

    # "Page 1 of N" pattern suggests document start
    if re.search(r'\bpage\s+1\s+of\s+\d+', lower[:400]):
        return True

    # Known document title in first 200 chars
    for starter in _DOC_STARTERS:
        if starter in first_200:
            return True

    return False


def _extract_pdf_text(pdf_path: str) -> str:
    """Extract text from PDF using pdfplumber (better layout fidelity) with PyMuPDF fallback."""
    try:
        with pdfplumber.open(pdf_path) as pdf:
            pages_text = []
            for page in pdf.pages[:10]:  # classify from first 10 pages
                t = page.extract_text() or ""
                pages_text.append(t)
            return "\n".join(pages_text)
    except Exception:
        pass

    # Fallback to PyMuPDF
    try:
        doc = fitz.open(pdf_path)
        text = "\n".join(doc[i].get_text() for i in range(min(10, len(doc))))
        doc.close()
        return text
    except Exception:
        return ""


def _extract_docx_text(path: str) -> str:
    try:
        from docx import Document
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs[:60])
    except Exception:
        return ""


def _extract_xlsx_text(path: str) -> str:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets[:2]:
            for row in ws.iter_rows(max_row=30, values_only=True):
                line = " ".join(str(c) for c in row if c is not None)
                if line.strip():
                    lines.append(line)
        return "\n".join(lines)
    except Exception:
        return ""


def _ocr_image_file(path: str) -> str:
    if not _OCR_AVAILABLE:
        return ""
    try:
        img = Image.open(path)
        return pytesseract.image_to_string(img)
    except Exception:
        return ""


def _ocr_image_bytes(img_bytes: bytes) -> str:
    if not _OCR_AVAILABLE:
        return ""
    try:
        img = Image.open(io.BytesIO(img_bytes))
        return pytesseract.image_to_string(img)
    except Exception:
        return ""
